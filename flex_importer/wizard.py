"""
Wizard de importación: parseo, auto-mapeo de columnas, limpieza y validación
previas a la creación del ImportJob.

El wizard nunca toca el pipeline de importación existente: su salida es un
archivo XLSX "normalizado" (encabezados = verbose_name de cada campo, valores
ya limpios y tipados) que entra por ImportProcessor como cualquier otro
archivo. Así no se requieren migraciones ni cambios en ImportJob.
"""
import csv
import io
import json
import re
import unicodedata
from datetime import datetime, date, timedelta
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


# Serial de Excel: días desde 1899-12-30. Rango plausible ~1900-2100.
EXCEL_EPOCH = datetime(1899, 12, 30)
EXCEL_SERIAL_MIN = 1
EXCEL_SERIAL_MAX = 80000

DATE_FORMATS = [
    '%d/%m/%Y %H:%M:%S', '%d/%m/%y %H:%M:%S',
    '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M',
    '%d/%m/%Y %H:%M', '%d/%m/%y %H:%M',
    '%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d',
    '%d-%m-%Y', '%d-%m-%y', '%Y/%m/%d',
]

TRUE_VALUES = {'true', 'yes', 'si', 'sí', '1', 't', 'x', 'verdadero'}
FALSE_VALUES = {'false', 'no', '0', 'f', '', 'falso'}

DEFAULT_CLEANING = {
    'recortar_espacios': True,      # strip a ambos extremos de todo texto
    'colapsar_espacios': False,     # runs de espacios internos -> uno solo
    'numeros_como_texto': True,     # 42.0 en campo texto -> "42"
    'coma_decimal': False,          # "1.234,56" -> 1234.56 (formato latino)
    'fechas_flexibles': True,       # multiples formatos + serial de Excel
    'omitir_filas_vacias': True,    # ya lo hace el lector, se expone por claridad
}


def detect_format(filename):
    """Deduce el formato por la extensión. Devuelve 'xlsx'|'csv'|'json'|None."""
    name = (filename or '').lower()
    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        return 'xlsx'
    if name.endswith('.csv') or name.endswith('.txt'):
        return 'csv'
    if name.endswith('.json'):
        return 'json'
    return None


def normalize_header(text):
    """Normaliza un encabezado para comparación: minúsculas, sin acentos,
    sin ' *' de requerido, espacios/underscores colapsados."""
    if text is None:
        return ''
    text = str(text).replace(' *', ' ').strip().lower()
    text = unicodedata.normalize('NFKD', text)
    text = ''.join(c for c in text if not unicodedata.combining(c))
    text = re.sub(r'[\s_]+', ' ', text)
    return text.strip()


class ParsedFile:
    """Resultado de leer un archivo: encabezados posicionales + filas."""

    def __init__(self, headers, rows, row_numbers):
        self.headers = headers          # list[str] (puede haber duplicados)
        self.rows = rows                # list[list] alineadas con headers
        self.row_numbers = row_numbers  # nro de fila original en el archivo

    @property
    def total(self):
        return len(self.rows)

    def duplicate_headers(self):
        seen, dups = set(), []
        for h in self.headers:
            key = normalize_header(h)
            if key and key in seen and h not in dups:
                dups.append(h)
            seen.add(key)
        return dups


def parse_file(file_obj, file_format, header_row=1, max_rows=None):
    """Lee el archivo completo a memoria como ParsedFile.

    A diferencia del processor, conserva las columnas por POSICIÓN, de modo
    que encabezados duplicados no pisan datos y el mapeo es por índice.
    `max_rows` acota el uso de memoria: si se supera, se aborta con ValueError.
    """
    if file_format == 'xlsx':
        return _parse_xlsx(file_obj, header_row, max_rows)
    if file_format == 'csv':
        return _parse_csv(file_obj, max_rows)
    if file_format == 'json':
        return _parse_json(file_obj, max_rows)
    raise ValueError(f'Formato no soportado: {file_format}')


def _parse_xlsx(file_obj, header_row, max_rows=None):
    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.active

    headers = []
    rows = []
    row_numbers = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if idx < header_row:
            continue
        if idx == header_row:
            # los encabezados vacíos intermedios se nombran "Columna N" para no
            # ocultar las columnas a su derecha (títulos con celdas combinadas)
            empty_idx = set()
            for col, cell in enumerate(row, start=1):
                if cell is None or str(cell).strip() == '':
                    empty_idx.add(len(headers))
                    headers.append(f'Columna {col}')
                else:
                    headers.append(str(cell).replace(' *', '').strip())
            while headers and (len(headers) - 1) in empty_idx:
                headers.pop()
            continue
        if max_rows is not None and len(rows) >= max_rows:
            wb.close()
            raise ValueError(f'El archivo supera el máximo de {max_rows} filas')
        values = list(row[:len(headers)])
        if all(v is None or str(v).strip() == '' for v in values):
            continue
        values += [None] * (len(headers) - len(values))
        rows.append(values)
        row_numbers.append(idx)
    wb.close()

    if not headers:
        raise ValueError('No se encontraron encabezados en el archivo')
    return ParsedFile(headers, rows, row_numbers)


def _decode_bytes(raw):
    for encoding in ('utf-8-sig', 'utf-8', 'cp1252', 'latin-1'):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError('No se pudo decodificar el archivo CSV (codificación desconocida)')


def _parse_csv(file_obj, max_rows=None):
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = _decode_bytes(raw)
    else:
        text = raw

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)
    all_rows = list(reader)
    if not all_rows:
        raise ValueError('El archivo CSV está vacío')

    headers = [str(h).replace(' *', '').strip() for h in all_rows[0]]
    while headers and headers[-1] == '':
        headers.pop()
    if not headers:
        raise ValueError('No se encontraron encabezados en el archivo')

    rows, row_numbers = [], []
    for idx, row in enumerate(all_rows[1:], start=2):
        values = [v if v != '' else None for v in row[:len(headers)]]
        if all(v is None or str(v).strip() == '' for v in values):
            continue
        if max_rows is not None and len(rows) >= max_rows:
            raise ValueError(f'El archivo supera el máximo de {max_rows} filas')
        values += [None] * (len(headers) - len(values))
        rows.append(values)
        row_numbers.append(idx)
    return ParsedFile(headers, rows, row_numbers)


def _parse_json(file_obj, max_rows=None):
    raw = file_obj.read()
    if isinstance(raw, bytes):
        raw = _decode_bytes(raw)
    data = json.loads(raw)

    if isinstance(data, dict) and 'data' in data:
        records = data['data']
    elif isinstance(data, list):
        records = data
    else:
        raise ValueError('JSON inválido: debe ser una lista o un objeto con propiedad "data"')
    if not isinstance(records, list):
        raise ValueError('JSON inválido: se esperaba una lista de objetos')
    for idx, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ValueError(f'JSON inválido: el registro {idx} no es un objeto')

    headers = []
    for record in records:
        for key in record.keys():
            if key not in headers:
                headers.append(str(key))
    if not headers:
        raise ValueError('El JSON no contiene registros con campos')

    rows, row_numbers = [], []
    for idx, record in enumerate(records, start=1):
        values = [record.get(h) for h in headers]
        if all(v is None or str(v).strip() == '' for v in values):
            continue
        if max_rows is not None and len(rows) >= max_rows:
            raise ValueError(f'El archivo supera el máximo de {max_rows} filas')
        rows.append(values)
        row_numbers.append(idx)
    return ParsedFile(headers, rows, row_numbers)


def public_field_info(importer_class):
    """get_field_info() sin el objeto Field (serializable a JSON)."""
    return [
        {
            'name': info['name'],
            'verbose_name': str(info['verbose_name']),
            'required': bool(info['required']),
            'type': info['type'],
        }
        for info in importer_class.get_field_info()
    ]


def duplicate_verbose_names(importer_class):
    """verbose_name repetidos: colisionarían como encabezados del normalizado."""
    seen, dups = set(), []
    for info in importer_class.get_field_info():
        v = str(info['verbose_name'])
        if v in seen and v not in dups:
            dups.append(v)
        seen.add(v)
    return dups


def auto_map(headers, field_info):
    """Propone mapping {field_name: índice de columna | None}.

    Prioridad por campo: verbose_name exacto > name exacto > verbose_name
    normalizado > name normalizado. Cada columna se asigna a lo sumo una vez.
    """
    norm_headers = [normalize_header(h) for h in headers]
    used = set()
    mapping = {}

    def find(predicate):
        for idx, header in enumerate(headers):
            if idx in used:
                continue
            if predicate(header, norm_headers[idx]):
                return idx
        return None

    for info in field_info:
        verbose = str(info['verbose_name'])
        name = info['name']
        idx = find(lambda h, nh: h == verbose)
        if idx is None:
            idx = find(lambda h, nh: h == name)
        if idx is None:
            idx = find(lambda h, nh: nh == normalize_header(verbose))
        if idx is None:
            idx = find(lambda h, nh: nh == normalize_header(name))
        mapping[info['name']] = idx
        if idx is not None:
            used.add(idx)
    return mapping


# ---------------------------------------------------------------------------
# Limpieza de valores
# ---------------------------------------------------------------------------

def _clean_text(value, ops):
    text = str(value)
    # caracteres de control que Excel no admite (romperían el XLSX normalizado)
    text = ILLEGAL_CHARACTERS_RE.sub('', text)
    if ops.get('recortar_espacios', True):
        text = text.strip()
    if ops.get('colapsar_espacios'):
        text = re.sub(r'\s+', ' ', text)
    return text


def _clean_number_string(text, ops):
    """Deja un string numérico listo para float()/Decimal()."""
    text = text.strip().replace(' ', '').replace(' ', '')
    # negativos contables: (1.234,56) o trailing minus "123-"
    negative = False
    if len(text) >= 2 and text.startswith('(') and text.endswith(')'):
        negative = True
        text = text[1:-1]
    if text.endswith('-'):
        negative = True
        text = text[:-1]
    # símbolos de moneda comunes en los archivos de la región
    text = re.sub(r'^[^0-9\-+.,]+|[^0-9.,]+$', '', text)
    if ops.get('coma_decimal'):
        text = text.replace('.', '').replace(',', '.')
    else:
        text = text.replace(',', '')
    if negative and text and not text.startswith('-'):
        text = '-' + text
    return text


def _to_naive(value):
    """Aware -> hora local del proyecto, naive (no descartar el offset)."""
    if value.tzinfo is None:
        return value
    from django.utils.timezone import get_current_timezone
    return value.astimezone(get_current_timezone()).replace(tzinfo=None)


def _parse_datetime(value, ops):
    if isinstance(value, datetime):
        return _to_naive(value)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if ops.get('fechas_flexibles', True) and EXCEL_SERIAL_MIN <= float(value) <= EXCEL_SERIAL_MAX:
            return EXCEL_EPOCH + timedelta(days=float(value))
        raise ValueError(f'"{value}" no parece una fecha')

    text = str(value).strip()
    if not text:
        raise ValueError('fecha vacía')
    try:
        return _to_naive(datetime.fromisoformat(text))
    except ValueError:
        pass
    if ops.get('fechas_flexibles', True):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        # serial de Excel llegado como texto
        try:
            serial = float(text)
            if EXCEL_SERIAL_MIN <= serial <= EXCEL_SERIAL_MAX:
                return EXCEL_EPOCH + timedelta(days=serial)
        except ValueError:
            pass
    raise ValueError(f'"{value}" no coincide con ningún formato de fecha conocido')


def clean_value(value, field_type, ops):
    """Limpia/convierte un valor crudo según el tipo del campo destino.

    Devuelve un valor apto para escribirse en el XLSX normalizado y que
    supera _convert_field_value() del importador. Lanza ValueError si el
    valor es irrecuperable.
    """
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        if stripped == '':
            return None

    if field_type in ('integer', 'float', 'decimal'):
        if isinstance(value, bool):
            raise ValueError(f'"{value}" no es un número')
        if isinstance(value, int) and field_type == 'integer':
            return value
        if isinstance(value, Decimal) and field_type == 'decimal':
            return value
        if isinstance(value, (int, float, Decimal)):
            number = float(value)
        else:
            text = _clean_number_string(str(value), ops)
            if not text or text in ('-', '+', '.'):
                raise ValueError(f'"{value}" no es un número')
            # preservar precisión exacta cuando el texto lo permite
            # (IDs largos, decimales con muchos dígitos)
            if field_type == 'integer':
                try:
                    return int(text)
                except ValueError:
                    pass
            if field_type == 'decimal':
                try:
                    return Decimal(text)
                except InvalidOperation:
                    pass
            try:
                number = float(text)
            except ValueError:
                raise ValueError(f'"{value}" no es un número válido')
        if field_type == 'integer':
            if abs(number - round(number)) > 1e-9:
                raise ValueError(f'"{value}" no es un número entero')
            return int(round(number))
        if field_type == 'decimal':
            try:
                return Decimal(str(number))
            except InvalidOperation:
                raise ValueError(f'"{value}" no es un decimal válido')
        return number

    if field_type in ('date', 'datetime'):
        parsed = _parse_datetime(value, ops)
        return parsed.date() if field_type == 'date' else parsed

    if field_type == 'boolean':
        if isinstance(value, bool):
            return value
        text = str(value).strip().lower()
        if text in TRUE_VALUES:
            return True
        if text in FALSE_VALUES:
            return False
        raise ValueError(f'"{value}" no es un valor booleano (si/no)')

    # text / email / default
    if ops.get('numeros_como_texto', True) and isinstance(value, float) and value.is_integer():
        return str(int(value))
    if not isinstance(value, str):
        value = str(value)
    return _clean_text(value, ops)


# ---------------------------------------------------------------------------
# Validación del dataset completo
# ---------------------------------------------------------------------------

def display_value(value, limit=60):
    if value is None:
        return ''
    text = str(value)
    return text if len(text) <= limit else text[:limit - 1] + '…'


def validate_dataset(parsed, mapping, importer_class, ops, preview_limit=15):
    """Aplica mapeo + limpieza a todas las filas y las valida contra el
    importador. Devuelve un dict serializable con el resultado.
    """
    field_info = importer_class.get_field_info()
    ops = {**DEFAULT_CLEANING, **(ops or {})}

    # instancia (no clase): los importadores pueden overridear validate_row
    # como método de instancia, igual que hace ImportProcessor
    importer_instance = importer_class()

    issues_by_field = {}
    invalid_rows = set()
    preview = []
    cleaned_rows = []          # list[dict field_name -> valor limpio]
    row_errors_list = []       # list[dict field_name -> mensaje], alineada con rows
    key_field = importer_class.get_key_field()
    key_values = {}

    for pos, values in enumerate(parsed.rows):
        row_number = parsed.row_numbers[pos]
        cleaned = {}
        row_errors = {}

        for info in field_info:
            name = info['name']
            col = mapping.get(name)
            raw = values[col] if col is not None and col < len(values) else None
            try:
                cleaned_value = clean_value(raw, info['type'], ops)
            except ValueError as exc:
                cleaned[name] = None
                row_errors[name] = str(exc)
                continue
            if info['required'] and cleaned_value is None:
                cleaned[name] = None
                row_errors[name] = 'Valor requerido ausente'
                continue
            cleaned[name] = cleaned_value

        # validación final con la lógica propia del importador; se pasan las
        # claves por name Y por verbose_name porque los overrides usan ambas
        if not row_errors:
            by_both = {}
            for info in field_info:
                by_both[info['name']] = cleaned[info['name']]
                by_both[str(info['verbose_name'])] = cleaned[info['name']]
            try:
                _, errors = importer_instance.validate_row(by_both)
            except Exception as exc:
                errors = [f'validate_row falló: {exc}']
            for err in errors:
                row_errors.setdefault('__row__', err)

        if row_errors:
            invalid_rows.add(row_number)
            for info in field_info:
                name = info['name']
                if name not in row_errors:
                    continue
                bucket = issues_by_field.setdefault(name, {
                    'campo': name,
                    'verbose': str(info['verbose_name']),
                    'errores': 0,
                    'ejemplos': [],
                })
                bucket['errores'] += 1
                if len(bucket['ejemplos']) < 5:
                    col = mapping.get(name)
                    raw = values[col] if col is not None and col < len(values) else None
                    bucket['ejemplos'].append({
                        'fila': row_number,
                        'valor': display_value(raw),
                        'error': row_errors[name],
                    })
            if '__row__' in row_errors:
                bucket = issues_by_field.setdefault('__row__', {
                    'campo': '__row__',
                    'verbose': 'Validación del importador',
                    'errores': 0,
                    'ejemplos': [],
                })
                bucket['errores'] += 1
                if len(bucket['ejemplos']) < 5:
                    bucket['ejemplos'].append({
                        'fila': row_number,
                        'valor': '',
                        'error': row_errors['__row__'],
                    })

        if key_field and not row_errors:
            key_val = cleaned.get(key_field)
            if key_val is not None:
                key_values.setdefault(key_val, []).append(row_number)

        cleaned_rows.append(cleaned)
        row_errors_list.append(row_errors)

        if len(preview) < preview_limit:
            preview.append({
                'fila': row_number,
                'valores': {
                    info['name']: display_value(cleaned.get(info['name']))
                    for info in field_info
                },
                'errores': {k: v for k, v in row_errors.items() if k != '__row__'},
                'error_fila': row_errors.get('__row__', ''),
            })

    duplicate_keys = {k: v for k, v in key_values.items() if len(v) > 1}

    return {
        'resultado': {
            'total': parsed.total,
            'validas': parsed.total - len(invalid_rows),
            'con_errores': len(invalid_rows),
            'por_campo': sorted(
                issues_by_field.values(), key=lambda b: -b['errores']
            ),
            'duplicados_clave': len(duplicate_keys),
            'campo_clave': key_field or '',
            'preview': preview,
        },
        'cleaned_rows': cleaned_rows,
        'row_errors_list': row_errors_list,
        'invalid_rows': invalid_rows,
    }


def _safe_cell(ws, row_idx, col_idx, value):
    """Escribe una celda neutralizando fórmulas ('=...' se fuerza a texto)."""
    cell = ws.cell(row=row_idx, column=col_idx, value=value)
    if isinstance(value, str) and value.startswith('='):
        cell.data_type = 's'
    return cell


def build_normalized_xlsx(parsed, cleaned_rows, row_errors_list, invalid_rows,
                          mapping, importer_class, skip_invalid=True,
                          header_row=1):
    """Genera el XLSX normalizado (encabezados = verbose_name, valores limpios).

    Si skip_invalid=False, en los campos que fallaron limpieza se escribe el
    valor CRUDO (como texto) para que el processor reporte el error real en
    vez de un falso "requerido ausente".

    El encabezado se escribe en `header_row` (los importadores con
    Meta.header_row > 1 releen el normalizado desde esa misma fila).
    Se agrega la columna `_fila_original` para que los errores del processor
    referencien la fila del archivo que subió el usuario.

    Devuelve (BytesIO, filas_incluidas, filas_omitidas).
    """
    field_info = importer_class.get_field_info()
    text_ops = dict(DEFAULT_CLEANING)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Datos'
    for col_idx, info in enumerate(field_info, start=1):
        _safe_cell(ws, header_row, col_idx, str(info['verbose_name']))
    _safe_cell(ws, header_row, len(field_info) + 1, '_fila_original')

    included = skipped = 0
    out_row = header_row + 1
    for pos, cleaned in enumerate(cleaned_rows):
        row_number = parsed.row_numbers[pos]
        row_errors = row_errors_list[pos] if pos < len(row_errors_list) else {}
        if skip_invalid and row_number in invalid_rows:
            skipped += 1
            continue

        row = []
        for info in field_info:
            name = info['name']
            value = cleaned.get(name)
            if value is None and name in row_errors:
                # rescatar el valor original que falló la limpieza
                col = mapping.get(name)
                raw = (parsed.rows[pos][col]
                       if col is not None and col < len(parsed.rows[pos]) else None)
                if raw is not None:
                    value = _clean_text(raw, text_ops)
            if isinstance(value, Decimal):
                value = float(value)
            row.append(value)

        if all(v is None or v == '' for v in row):
            skipped += 1
            continue
        for col_idx, value in enumerate(row, start=1):
            _safe_cell(ws, out_row, col_idx, value)
        _safe_cell(ws, out_row, len(field_info) + 1, row_number)
        out_row += 1
        included += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, included, skipped
